"""Модуль для экспорта данных в Excel и PDF."""
import openpyxl
from openpyxl.styles import Font
from fpdf import FPDF

# Словарь для транслитерации (замена русских букв на английские)
TRANSLIT_DICT = {
    'а': 'a', 'б': 'b', 'в': 'v', 'г': 'g', 'д': 'd', 'е': 'e', 'ё': 'e',
    'ж': 'zh', 'з': 'z', 'и': 'i', 'й': 'y', 'к': 'k', 'л': 'l', 'м': 'm',
    'н': 'n', 'о': 'o', 'п': 'p', 'р': 'r', 'с': 's', 'т': 't', 'у': 'u',
    'ф': 'f', 'х': 'kh', 'ц': 'ts', 'ч': 'ch', 'ш': 'sh', 'щ': 'shch',
    'ъ': '', 'ы': 'y', 'ь': '', 'э': 'e', 'ю': 'yu', 'я': 'ya',
    'А': 'A', 'Б': 'B', 'В': 'V', 'Г': 'G', 'Д': 'D', 'Е': 'E', 'Ё': 'E',
    'Ж': 'Zh', 'З': 'Z', 'И': 'I', 'Й': 'Y', 'К': 'K', 'Л': 'L', 'М': 'M',
    'Н': 'N', 'О': 'O', 'П': 'P', 'Р': 'R', 'С': 'S', 'Т': 'T', 'У': 'U',
    'Ф': 'F', 'Х': 'Kh', 'Ц': 'Ts', 'Ч': 'Ch', 'Ш': 'Sh', 'Щ': 'Shch',
    'Ъ': '', 'Ы': 'Y', 'Ь': '', 'Э': 'E', 'Ю': 'Yu', 'Я': 'Ya'
}

def translit(text: str) -> str:
    """Переводит русские буквы в английские (транслит)."""
    return "".join(TRANSLIT_DICT.get(char, char) for char in str(text))


# --- ЭКСПОРТ В EXCEL ---
def export_to_excel(entries_data: list, filepath: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Мои пароли"
    headers = ["Название", "Логин", "Пароль", "Надёжность", "Заметки", "Обновлено"]
    ws.append(headers)
    for cell in ws[1]: 
        cell.font = Font(bold=True)
    for row in entries_data:
        ws.append([row["title"], row["username"], row["password"], row["strength"], row["notes"], row["date"]])
    for col in ws.columns:
        max_length = 0
        for cell in col:
            try:
                if len(str(cell.value)) > max_length: 
                    max_length = len(str(cell.value))
            except: 
                pass
        ws.column_dimensions[col[0].column_letter].width = max_length + 2
    wb.save(filepath)


# --- ЭКСПОРТ В PDF ---
class PDF(FPDF):
    def header(self):
        self.set_font('helvetica', 'B', 16)
        self.cell(0, 10, 'Password Manager Export', border=0, align='C')
        self.ln(15)

    def footer(self):
        self.set_y(-15)
        self.set_font('helvetica', '', 8)
        self.cell(0, 10, f'Page {self.page_no()}', 0, 0, 'C')

def export_to_pdf(entries_data: list, filepath: str) -> None:
    pdf = PDF(orientation='L', unit='mm', format='A4')
    pdf.add_page()
    
    # Используем стандартный встроенный шрифт (без файлов)
    pdf.set_font('helvetica', 'B', 11)
    pdf.set_fill_color(230, 230, 230)
    
    w = [50, 65, 75, 40, 47]
    
    pdf.cell(w[0], 10, 'Title', border=1, align='C', fill=True)
    pdf.cell(w[1], 10, 'Login', border=1, align='C', fill=True)
    pdf.cell(w[2], 10, 'Password', border=1, align='C', fill=True)
    pdf.cell(w[3], 10, 'Strength', border=1, align='C', fill=True)
    pdf.cell(w[4], 10, 'Updated', border=1, align='C', fill=True)
    pdf.ln()

    pdf.set_font('helvetica', '', 10)
    for row in entries_data:
        # Переводим все данные в транслит
        title = translit(row["title"])
        if len(title) > 25: 
            title = title[:22] + "..."
            
        username = translit(row["username"])
        password = translit(row["password"])
        strength = translit(row["strength"])
        date = translit(row["date"])
        
        pdf.cell(w[0], 10, title, border=1)
        pdf.cell(w[1], 10, username, border=1)
        pdf.cell(w[2], 10, password, border=1)
        pdf.cell(w[3], 10, strength, border=1, align='C')
        pdf.cell(w[4], 10, date, border=1, align='C')
        pdf.ln()

    pdf.output(filepath)